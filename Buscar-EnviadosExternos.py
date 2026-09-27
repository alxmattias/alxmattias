#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Buscar-EnviadosExternos.py
===========================
Script dedicado: busca los correos que un usuario ENVIÓ hacia dominios
EXTERNOS (omite los dominios de DOMINIOS_INTERNOS), directamente sobre los
CSV de Message Trace, y genera un reporte Excel (+ PDF si Excel/pywin32
están disponibles).

USO:
    python Buscar-EnviadosExternos.py --usuario correo@ine.mx
    python Buscar-EnviadosExternos.py --usuario correo@ine.mx --desde 2024-01-01 --hasta 2026-09-25
    python Buscar-EnviadosExternos.py                     (modo interactivo, sin argumentos)

--desde / --hasta son OPCIONALES (formato AAAA-MM-DD, horario CDMX). Si no
se indican, se revisa TODO el histórico disponible. Cuando SÍ se indican,
el script descarta de entrada los archivos MessageTrace_AAAAMMDD-HHMMSS.csv
cuyo nombre cae fuera del rango, para no tener que abrirlos.

Requisitos (una sola vez):
    pip install openpyxl python-dateutil tzdata pywin32
"""

import argparse
import concurrent.futures
import csv
import glob
import os
import re
import sys
import time
from datetime import datetime
from zoneinfo import ZoneInfo

csv.field_size_limit(min(2**31 - 1, sys.maxsize))

try:
    from dateutil import parser as dtparser
    HAVE_DATEUTIL = True
except ImportError:
    HAVE_DATEUTIL = False

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

# ============================================================
# CONFIGURACION
# ============================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_DIR = r"\\VM-ADMON-WIN\E$\Bitacoras"
OUT_DIR = SCRIPT_DIR

TZ_CDMX = ZoneInfo("America/Mexico_City")
TZ_UTC = ZoneInfo("UTC")

EXCLUDE_FILES = {"fechas.csv", "usuarios.csv", "refechas.csv"}

# Dominios INTERNOS: cualquier correo enviado hacia uno de estos dominios se
# OMITE del reporte (solo interesa lo enviado hacia fuera de la organización).
DOMINIOS_INTERNOS = {
    "ine.mx",
    "sec.ife.org.mx",
    "ife.org.mx",
    "listas.ine.mx",
    "inemexico.mail.onmicrosoft.com",
}

STATUS_INFO = {
    "delivered": ("Entregado", "El mensaje llegó correctamente al buzón del destinatario."),
    "failed": ("Error", "El mensaje no pudo entregarse; fue rechazado o rebotado."),
    "pending": ("Pendiente", "El mensaje todavía está en proceso de entrega."),
    "resolved": ("Resuelto", "Notificación de un sistema de monitoreo que ya fue atendida."),
    "quarantined": ("En cuarentena", "El sistema de seguridad detuvo el mensaje por sospecha de spam o virus."),
    "filteredasspam": ("Filtrado como spam", "Se identificó como correo no deseado y no llegó a la bandeja principal."),
    "gettingstatus": ("Consultando estatus", "El sistema todavía está verificando qué pasó con el mensaje."),
    "expanded": ("Expandido", "El mensaje se envió a una lista o grupo de correo y se distribuyó a sus integrantes."),
    "recalled": ("Recall solicitado", "El remitente intentó recuperar (cancelar) este mensaje después de haberlo enviado."),
    "removedbymessagerecall": ("Retirado por recall", "El mensaje fue eliminado del buzón del destinatario porque el remitente lo recuperó a tiempo."),
    "poison": ("Mensaje dañado", "El sistema detectó un problema grave en el mensaje y lo detuvo por seguridad."),
    "none": ("Sin estatus", "No hay información de estatus registrada para este evento."),
}


def status_label(status_raw):
    info = STATUS_INFO.get((status_raw or "").strip().lower())
    return info[0] if info else (status_raw or "Sin estatus")


def status_desc(status_raw):
    info = STATUS_INFO.get((status_raw or "").strip().lower())
    return info[1] if info else "Estatus no documentado; consulte con el área de soporte si tiene dudas."


LILA_INST = "674092"
LILA_OSCURO = "49276F"
GRIS_CLARO = "F2F2F2"
BLANCO = "FFFFFF"
VERDE_OK = "E2EFDA"
ROJO_ERR = "FCE4E4"


def F(**k):
    return Font(name="Arial", **k)


# ============================================================
# FECHAS
# ============================================================
def normaliza_fecha(texto):
    t = texto.replace("\u00a0", " ")
    t = re.sub(r"a\.\s*m\.", "AM", t, flags=re.IGNORECASE)
    t = re.sub(r"p\.\s*m\.", "PM", t, flags=re.IGNORECASE)
    return t.strip()


def parsea_fecha_cdmx(texto):
    t = normaliza_fecha(texto)
    dt_naive = None
    if HAVE_DATEUTIL:
        try:
            dt_naive = dtparser.parse(t, dayfirst=True)
        except (ValueError, OverflowError, TypeError):
            dt_naive = None
    if dt_naive is None:
        for fmt in ("%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %I:%M %p", "%d/%m/%Y %H:%M"):
            try:
                dt_naive = datetime.strptime(t, fmt)
                break
            except ValueError:
                continue
    if dt_naive is None:
        return None
    return dt_naive.replace(tzinfo=TZ_CDMX)


def parsea_fecha_arg(texto):
    try:
        return datetime.strptime(texto.strip(), "%Y-%m-%d").date()
    except ValueError:
        raise argparse.ArgumentTypeError(
            f"Fecha inválida: '{texto}' (usa el formato AAAA-MM-DD, por ejemplo 2026-06-01)"
        )


RE_FECHA_ARCHIVO = re.compile(r"MessageTrace_(\d{8})")


def extrae_fecha_archivo(nombre):
    m = RE_FECHA_ARCHIVO.match(nombre)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%Y%m%d").date()
    except ValueError:
        return None


def archivos_todos(desde=None, hasta=None):
    """
    Todos los MessageTrace_*.csv disponibles. Si se dan desde/hasta,
    descarta de entrada (por nombre de archivo, MessageTrace_AAAAMMDD-...)
    los que caen fuera del rango, sin necesidad de abrirlos. Un archivo
    cuyo nombre no trae fecha reconocible siempre se incluye.
    """
    candidatos = []
    for f in glob.glob(os.path.join(CSV_DIR, "MessageTrace_*.csv")):
        nombre = os.path.basename(f)
        if nombre in EXCLUDE_FILES:
            continue
        if desde is not None or hasta is not None:
            fecha_archivo = extrae_fecha_archivo(nombre)
            if fecha_archivo is not None:
                if desde is not None and fecha_archivo < desde:
                    continue
                if hasta is not None and fecha_archivo > hasta:
                    continue
        candidatos.append(f)
    return sorted(candidatos)


def lee_csv_robusto(path):
    for enc in ("utf-8", "cp1252"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                return list(csv.reader(fh))
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="cp1252", errors="replace", newline="") as fh:
        return list(csv.reader(fh))


def formatea_duracion(segundos):
    segundos = max(0, int(round(segundos)))
    h, resto = divmod(segundos, 3600)
    m, s = divmod(resto, 60)
    if h:
        return f"{h}h {m:02d}m"
    if m:
        return f"{m}m {s:02d}s"
    return f"{s}s"


def imprime_barra_progreso(actual, total, tiempo_inicio, ancho=30):
    if total <= 0:
        return
    pct = actual / total
    llenado = int(ancho * pct)
    barra = "#" * llenado + "-" * (ancho - llenado)
    transcurrido = time.time() - tiempo_inicio
    if actual > 0:
        eta_txt = formatea_duracion((transcurrido / actual) * (total - actual)) if actual < total else "0s"
    else:
        eta_txt = "calculando..."
    sys.stdout.write(
        f"\r  [{barra}] {actual}/{total} archivos ({pct * 100:5.1f}%)"
        f" | transcurrido: {formatea_duracion(transcurrido)}"
        f" | restante (estimado): {eta_txt}   "
    )
    sys.stdout.flush()
    if actual >= total:
        sys.stdout.write("\n")
        sys.stdout.flush()


# ============================================================
# BÚSQUEDA
# ============================================================
def _procesa_un_archivo(path, usuario, desde, hasta):
    nombre_archivo = os.path.basename(path)
    try:
        filas = lee_csv_robusto(path)
    except Exception as e:
        return {"archivo": nombre_archivo, "error": str(e), "omitido_estructura": None, "matches": []}

    if not filas:
        return {"archivo": nombre_archivo, "error": None, "omitido_estructura": None, "matches": []}

    if len(filas[0]) != 8:
        return {
            "archivo": nombre_archivo,
            "error": None,
            "omitido_estructura": f"{nombre_archivo} ({len(filas[0])} columnas, se esperaban 8)",
            "matches": [],
        }

    matches = []
    for fila in filas:
        if len(fila) != 8:
            continue
        received, sender, recipient, subject, status, fromip, size, msgid = fila[:8]
        if received.strip().lower() == "received":
            continue
        s_low, r_low = sender.strip().lower(), recipient.strip().lower()
        if s_low != usuario:
            continue
        dominio_destino = r_low.rsplit("@", 1)[-1] if "@" in r_low else ""
        if dominio_destino in DOMINIOS_INTERNOS:
            continue

        dt_cdmx = parsea_fecha_cdmx(received)
        if dt_cdmx is None:
            continue
        if desde is not None and dt_cdmx.date() < desde:
            continue
        if hasta is not None and dt_cdmx.date() > hasta:
            continue

        matches.append({
            "fecha_original": received.strip(),
            "fecha_cdmx": dt_cdmx,
            "fecha_utc": dt_cdmx.astimezone(TZ_UTC),
            "sender": sender,
            "recipient": recipient,
            "subject": subject,
            "status": status,
            "size": size,
            "msgid": msgid,
            "_archivo_origen": nombre_archivo,
        })

    return {"archivo": nombre_archivo, "error": None, "omitido_estructura": None, "matches": matches}


def busca(usuario, desde=None, hasta=None, progreso=True, hilos=8):
    usuario = usuario.strip().lower()
    archivos = archivos_todos(desde=desde, hasta=hasta)
    total_archivos = len(archivos)
    if progreso:
        rango_txt = (
            f"del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}"
            if (desde or hasta) else "TODO el histórico disponible"
        )
        print(f"Revisando {total_archivos} archivo(s) — {rango_txt} — enviados por {usuario} a externos...")

    resultados = []
    vistos = {}
    duplicados = 0
    ejemplos_duplicados = []
    MAX_EJEMPLOS_DUPLICADOS = 20
    archivos_omitidos_estructura = []
    archivos_procesados_ok = 0
    archivos_con_error = []
    tiempo_inicio_busqueda = time.time()

    completados = 0
    with concurrent.futures.ThreadPoolExecutor(max_workers=hilos) as executor:
        futuros = {
            executor.submit(_procesa_un_archivo, path, usuario, desde, hasta): path
            for path in archivos
        }
        for futuro in concurrent.futures.as_completed(futuros):
            completados += 1
            if progreso:
                imprime_barra_progreso(completados, total_archivos, tiempo_inicio_busqueda)
            r = futuro.result()
            if r["error"]:
                archivos_con_error.append(f"{r['archivo']}: {r['error']}")
                continue
            if r["omitido_estructura"]:
                archivos_omitidos_estructura.append(r["omitido_estructura"])
                continue
            archivos_procesados_ok += 1

            for x in r["matches"]:
                s_low = x["sender"].strip().lower()
                r_low = x["recipient"].strip().lower()
                huella = (x["msgid"].strip(), x["fecha_original"], s_low, r_low, x["status"].strip().lower())
                if huella in vistos:
                    duplicados += 1
                    if len(ejemplos_duplicados) < MAX_EJEMPLOS_DUPLICADOS:
                        archivo_previo, subject_previo, size_previo = vistos[huella]
                        ejemplos_duplicados.append({
                            "msgid": x["msgid"].strip(),
                            "received": x["fecha_original"],
                            "sender": x["sender"],
                            "recipient": x["recipient"],
                            "status": x["status"],
                            "archivo_1": archivo_previo,
                            "subject_1": subject_previo,
                            "size_1": size_previo,
                            "archivo_2": x["_archivo_origen"],
                            "subject_2": x["subject"],
                            "size_2": x["size"],
                        })
                    continue
                vistos[huella] = (x["_archivo_origen"], x["subject"], x["size"])
                resultados.append(x)

    resultados.sort(key=lambda r: r["fecha_cdmx"])

    auditoria = {
        "archivos_candidatos": total_archivos,
        "archivos_procesados_ok": archivos_procesados_ok,
        "archivos_omitidos_estructura": archivos_omitidos_estructura,
        "archivos_con_error": archivos_con_error,
        "duplicados": duplicados,
        "ejemplos_duplicados": ejemplos_duplicados,
        "tiempo_busqueda_seg": time.time() - tiempo_inicio_busqueda,
    }
    return resultados, auditoria


# ============================================================
# EXCEL
# ============================================================
def oculta_no_usado(ws, last_col, last_row, col_buffer=40, row_buffer=300):
    for c in range(last_col + 1, last_col + 1 + col_buffer):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    for rr in range(last_row + 1, last_row + 1 + row_buffer):
        ws.row_dimensions[rr].hidden = True


def crea_hoja_detalle(ws, lista_datos):
    thin = Side(style="thin", color="BFBFBF")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_bottom = Border(bottom=Side(style="thin", color="BFBFBF"))
    font_detalle = F(size=8)

    ws.sheet_view.showGridLines = False
    ancho_num = max(4, len(str(len(lista_datos))) + 2)
    widths = [ancho_num, 18, 18, 26, 26, 45, 14, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    LAST_COL = 8
    r = 1
    c = ws.cell(r, 2, "DETALLE DE CORREOS ENVIADOS A DOMINIOS EXTERNOS")
    c.font = F(bold=True, size=14, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=LILA_INST)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LAST_COL)
    ws.row_dimensions[r].height = 24
    r += 2

    headers = ["#", "FECHA ORIGINAL", "FECHA (CDMX / UTC-6)", "REMITENTE", "DESTINATARIO", "ASUNTO", "ESTATUS", "TAMAÑO (KB)"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(r, j, h)
        c.font = F(bold=True, size=9, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=LILA_OSCURO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_all
    ws.row_dimensions[r].height = 22
    r += 1

    tabla_ini = r
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    fill_zebra = PatternFill("solid", fgColor=GRIS_CLARO)

    for idx, x in enumerate(lista_datos, start=1):
        st_es = status_label(x["status"])
        try:
            kb = round(int(str(x["size"]).strip() or 0) / 1024, 1)
        except (ValueError, TypeError):
            kb = ""
        vals = [idx, x["fecha_original"], x["fecha_cdmx"].strftime("%d/%m/%Y %H:%M:%S"),
                x["sender"], x["recipient"], x["subject"], st_es, kb]
        zebra = idx % 2 == 0
        for j, v in enumerate(vals, start=1):
            c = ws.cell(r, j, v)
            c.font = font_detalle
            c.alignment = align_center if j in (1, 2, 3, 7, 8) else align_left
            c.border = border_bottom
            if zebra:
                c.fill = fill_zebra
        r += 1

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.print_area = f"A1:{get_column_letter(LAST_COL)}{r - 1}"
    ws.print_title_rows = f"{tabla_ini - 1}:{tabla_ini - 1}"
    ws.freeze_panes = ws.cell(tabla_ini, 1)
    oculta_no_usado(ws, LAST_COL, r - 1)


def genera_excel(resultados, usuario, desde, hasta, out_path):
    thin = Side(style="thin", color="BFBFBF")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Resumen"
    ws1.sheet_view.showGridLines = False

    LAST_COL = 7
    widths = [6, 20, 26, 26, 45, 14, 10]
    for i, w in enumerate(widths, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    r = 1
    c = ws1.cell(r, 2, "REPORTE DE CORREOS ENVIADOS A DOMINIOS EXTERNOS")
    c.font = F(bold=True, size=14, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=LILA_INST)
    c.alignment = Alignment(horizontal="center")
    ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LAST_COL)
    ws1.row_dimensions[r].height = 24
    r += 2

    def bloque_titulo(row, texto):
        ws1.merge_cells(start_row=row, start_column=2, end_row=row, end_column=LAST_COL)
        c = ws1.cell(row, 2, texto)
        c.font = F(bold=True, size=11, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=LILA_OSCURO)
        c.alignment = Alignment(horizontal="left", indent=1, vertical="center")
        ws1.row_dimensions[row].height = 20

    bloque_titulo(r, "DATOS DE LA CONSULTA")
    r += 1
    if resultados:
        primer_fecha = min(x["fecha_cdmx"] for x in resultados).strftime("%d/%m/%Y")
        ultima_fecha = max(x["fecha_cdmx"] for x in resultados).strftime("%d/%m/%Y")
        rango_encontrado_txt = f"{primer_fecha} al {ultima_fecha} (horario CDMX)"
    else:
        rango_encontrado_txt = "Sin resultados en el rango/histórico revisado"
    rango_pedido_txt = (
        f"{desde.strftime('%d/%m/%Y') if desde else 'inicio del histórico'} al "
        f"{hasta.strftime('%d/%m/%Y') if hasta else 'hoy'}"
        if (desde or hasta) else "TODO el histórico disponible (sin filtro de fecha)"
    )
    datos = [
        ("Usuario consultado:", usuario),
        ("Modo de búsqueda:", "Correos ENVIADOS por el usuario hacia dominios externos"),
        ("Dominios internos omitidos:", ", ".join(sorted(DOMINIOS_INTERNOS))),
        ("Rango solicitado:", rango_pedido_txt),
        ("Rango real encontrado:", rango_encontrado_txt),
        ("Fecha de generación:", datetime.now(TZ_CDMX).strftime("%d/%m/%Y %H:%M hrs (CDMX)")),
    ]
    for etiqueta, valor in datos:
        c1 = ws1.cell(r, 2, etiqueta)
        c1.font = F(bold=True, size=10)
        ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws1.merge_cells(start_row=r, start_column=4, end_row=r, end_column=LAST_COL)
        ws1.cell(r, 4, valor).font = F(size=10)
        r += 1
    r += 1

    por_status = {}
    for x in resultados:
        st = status_label(x["status"])
        por_status[st] = por_status.get(st, 0) + 1

    bloque_titulo(r, "RESUMEN CONSOLIDADO")
    r += 1
    c1 = ws1.cell(r, 2, "Total de correos enviados a externos")
    c1.font = F(bold=True, size=10)
    ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws1.merge_cells(start_row=r, start_column=6, end_row=r, end_column=LAST_COL)
    c2 = ws1.cell(r, 6, len(resultados))
    c2.font = F(size=13, bold=True, color=LILA_INST)
    c2.alignment = Alignment(horizontal="center")
    for col in range(2, LAST_COL + 1):
        ws1.cell(r, col).border = border_all
    ws1.row_dimensions[r].height = 20
    r += 2

    if por_status:
        ws1.cell(r, 2, "Desglose por estatus").font = F(bold=True, size=10, color=LILA_OSCURO)
        r += 1
        hdr = ws1.cell(r, 2, "Estatus")
        hdr.font = F(bold=True, size=9, color=BLANCO)
        hdr.fill = PatternFill("solid", fgColor=LILA_INST)
        ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        hdr2 = ws1.cell(r, 6, "Cantidad")
        hdr2.font = F(bold=True, size=9, color=BLANCO)
        hdr2.fill = PatternFill("solid", fgColor=LILA_INST)
        ws1.merge_cells(start_row=r, start_column=6, end_row=r, end_column=LAST_COL)
        for col in range(2, LAST_COL + 1):
            ws1.cell(r, col).border = border_all
        r += 1
        for st, cnt in sorted(por_status.items(), key=lambda x: -x[1]):
            ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            ws1.cell(r, 2, st).font = F(size=9)
            ws1.merge_cells(start_row=r, start_column=6, end_row=r, end_column=LAST_COL)
            c2 = ws1.cell(r, 6, cnt)
            c2.font = F(size=9)
            c2.alignment = Alignment(horizontal="center")
            low = st.lower()
            fill = VERDE_OK if "entreg" in low else (ROJO_ERR if "error" in low else GRIS_CLARO)
            for col in range(2, LAST_COL + 1):
                ws1.cell(r, col).fill = PatternFill("solid", fgColor=fill)
                ws1.cell(r, col).border = border_all
            r += 1
        r += 1

        raws_presentes = {}
        for x in resultados:
            raws_presentes[status_label(x["status"])] = x["status"]
        bloque_titulo(r, "NOMENCLATURA DE ESTATUS")
        r += 1
        for st_label in sorted(raws_presentes.keys()):
            desc = status_desc(raws_presentes[st_label])
            c1 = ws1.cell(r, 2, st_label)
            c1.font = F(bold=True, size=9, color=LILA_OSCURO)
            ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            c1.alignment = Alignment(vertical="top")
            c2 = ws1.cell(r, 4, desc)
            c2.font = F(size=9)
            ws1.merge_cells(start_row=r, start_column=4, end_row=r, end_column=LAST_COL)
            c2.alignment = Alignment(vertical="top", wrap_text=True)
            for col in range(2, LAST_COL + 1):
                ws1.cell(r, col).border = border_all
            ws1.row_dimensions[r].height = 26
            r += 1

    ws1.page_setup.orientation = "landscape"
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 0
    ws1.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws1.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws1.print_area = f"A1:{get_column_letter(LAST_COL)}{r}"

    ws2 = wb.create_sheet(title="Enviados a externos")
    crea_hoja_detalle(ws2, resultados)

    try:
        wb.save(out_path)
    except PermissionError:
        base, ext = os.path.splitext(out_path)
        alterno = f"{base}_{datetime.now(TZ_CDMX).strftime('%H%M%S')}{ext}"
        print(f"[AVISO] No se pudo guardar en {out_path} (¿está abierto en Excel?). Se guardó como: {alterno}")
        wb.save(alterno)
        out_path = alterno
    return len(resultados), out_path


def escribe_log_auditoria(usuario, auditoria, total_encontrados, xlsx_path):
    log_dir = os.path.join(OUT_DIR, "Logs")
    os.makedirs(log_dir, exist_ok=True)
    safe_user = re.sub(r"[^a-zA-Z0-9]", "_", usuario.strip().lower())
    log_path = os.path.join(log_dir, f"{safe_user}.log")

    omitidos = auditoria["archivos_omitidos_estructura"]
    con_error = auditoria.get("archivos_con_error") or []

    lineas = []
    lineas.append("=" * 78)
    lineas.append(f"Auditoría de búsqueda — {datetime.now(TZ_CDMX).strftime('%d/%m/%Y %H:%M:%S hrs (CDMX)')}")
    lineas.append(f"Usuario consultado : {usuario}")
    lineas.append("Modo de búsqueda   : Correos ENVIADOS por el usuario hacia dominios externos"
                  f" (internos excluidos: {', '.join(sorted(DOMINIOS_INTERNOS))})")
    lineas.append("-" * 78)
    lineas.append(f"Archivos fuente revisados : {auditoria['archivos_procesados_ok']} de {auditoria['archivos_candidatos']} candidatos")
    lineas.append(f"Archivos con estructura no estándar (omitidos) : {len(omitidos) if omitidos else 'Ninguno'}")
    for nombre in omitidos:
        lineas.append(f"    - {nombre}")
    lineas.append(f"Archivos con error de lectura : {len(con_error) if con_error else 'Ninguno'}")
    for nombre in con_error:
        lineas.append(f"    - {nombre}")
    lineas.append(f"Registros duplicados detectados y omitidos : {auditoria['duplicados']}")
    lineas.append(f"Tiempo de búsqueda : {formatea_duracion(auditoria['tiempo_busqueda_seg'])}")
    lineas.append("-" * 78)
    lineas.append(f"Mensajes únicos encontrados : {total_encontrados}")
    lineas.append(f"Reporte generado : {xlsx_path}")
    lineas.append("=" * 78)
    lineas.append("")

    with open(log_path, "a", encoding="utf-8") as fh:
        fh.write("\n".join(lineas) + "\n")
    return log_path


def exporta_pdf(xlsx_path, pdf_path):
    try:
        import win32com.client
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(xlsx_path)
        wb.Sheets.Select()
        wb.ActiveSheet.ExportAsFixedFormat(0, pdf_path)
        wb.Close(False)
        excel.Quit()
        return True
    except Exception as e:
        print(f"[AVISO] No se pudo generar el PDF automáticamente: {e}")
        print("        Abre el .xlsx en Excel y usa 'Guardar como > PDF' manualmente.")
        return False


# ============================================================
# MAIN
# ============================================================
def main():
    global CSV_DIR
    tiempo_inicio_total = time.time()
    ap = argparse.ArgumentParser(
        description="Busca los correos que un usuario envió hacia dominios EXTERNOS (omite DOMINIOS_INTERNOS)."
    )
    ap.add_argument("--usuario", help="Correo del usuario a buscar")
    ap.add_argument("--desde", type=parsea_fecha_arg, default=None, help="Fecha inicial AAAA-MM-DD (opcional).")
    ap.add_argument("--hasta", type=parsea_fecha_arg, default=None, help="Fecha final AAAA-MM-DD (opcional).")
    ap.add_argument("--csv-dir", dest="csv_dir", default=None, help=f"Carpeta con los MessageTrace_*.csv (default: {CSV_DIR}).")
    ap.add_argument("--hilos", type=int, default=8, help="Archivos a procesar en paralelo (default 8).")
    args = ap.parse_args()

    if args.csv_dir:
        CSV_DIR = args.csv_dir
    print(f"Leyendo bitácoras desde: {CSV_DIR}")

    if args.desde and args.hasta and args.desde > args.hasta:
        print("[ERROR] --desde no puede ser posterior a --hasta.")
        sys.exit(1)

    usuario = args.usuario or input("Usuario (correo): ").strip()

    print(
        f"\nBúsqueda de enviados a externos para {usuario}"
        + (f" del {args.desde.strftime('%d/%m/%Y')} al {args.hasta.strftime('%d/%m/%Y')}"
           if (args.desde or args.hasta) else " — TODO el histórico disponible")
        + "..."
    )
    resultados, auditoria = busca(usuario, desde=args.desde, hasta=args.hasta, hilos=args.hilos)
    print(f"Encontrados: {len(resultados)} mensaje(s) único(s).")
    print(f"Archivos revisados: {auditoria['archivos_procesados_ok']} de {auditoria['archivos_candidatos']} candidatos.")
    if auditoria["archivos_omitidos_estructura"]:
        print(f"[ATENCIÓN] {len(auditoria['archivos_omitidos_estructura'])} archivo(s) con estructura distinta a la esperada:")
        for nombre in auditoria["archivos_omitidos_estructura"]:
            print(f"    - {nombre}")
    if auditoria.get("archivos_con_error"):
        print(f"[ATENCIÓN] {len(auditoria['archivos_con_error'])} archivo(s) con error de lectura:")
        for nombre in auditoria["archivos_con_error"]:
            print(f"    - {nombre}")
    if auditoria["duplicados"]:
        print(f"Se detectaron y omitieron {auditoria['duplicados']} registro(s) duplicado(s) entre los archivos.")
    print(f"Tiempo de búsqueda: {formatea_duracion(auditoria['tiempo_busqueda_seg'])}")

    if not resultados:
        print("No se generará reporte: no hubo resultados para ese usuario/rango.")
        log_path = escribe_log_auditoria(usuario, auditoria, 0, "N/A (sin resultados, no se generó reporte)")
        print(f"Auditoría de la búsqueda agregada al log: {log_path}")
        print(f"Tiempo total de ejecución: {formatea_duracion(time.time() - tiempo_inicio_total)}")
        sys.exit(0)

    safe_user = re.sub(r"[^a-zA-Z0-9]", "_", usuario)
    sufijo_fechas = (
        f"_{args.desde.strftime('%Y%m%d') if args.desde else 'inicio'}-{args.hasta.strftime('%Y%m%d') if args.hasta else 'hoy'}"
        if (args.desde or args.hasta) else "_TODO"
    )
    nombre_base = f"EnviadosExternos_{safe_user}{sufijo_fechas}"
    xlsx_path = os.path.join(OUT_DIR, nombre_base + ".xlsx")

    tiempo_inicio_excel = time.time()
    total, xlsx_path = genera_excel(resultados, usuario, args.desde, args.hasta, xlsx_path)
    print(f"Excel generado ({formatea_duracion(time.time() - tiempo_inicio_excel)}): {xlsx_path}")

    log_path = escribe_log_auditoria(usuario, auditoria, len(resultados), xlsx_path)
    print(f"Auditoría de la búsqueda agregada al log: {log_path}")

    pdf_path = os.path.splitext(xlsx_path)[0] + ".pdf"
    tiempo_inicio_pdf = time.time()
    if exporta_pdf(xlsx_path, pdf_path):
        print(f"PDF generado ({formatea_duracion(time.time() - tiempo_inicio_pdf)}): {pdf_path}")

    print(f"Tiempo total de ejecución: {formatea_duracion(time.time() - tiempo_inicio_total)}")


if __name__ == "__main__":
    main()
